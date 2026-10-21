import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEAL = RGBColor(0x0D, 0x94, 0x88)
DARK = RGBColor(0x1E, 0x29, 0x3B)

doc = Document()

# Base style
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.font.color.rgb = DARK

# Heading colours
for lvl in (1, 2, 3):
    h = doc.styles[f'Heading {lvl}']
    h.font.color.rgb = TEAL if lvl == 1 else DARK
    h.font.name = 'Calibri'

def add_rich(paragraph, text):
    """Add text with **bold** inline parsing."""
    for part in re.split(r'(\*\*.*?\*\*)', text):
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part:
            paragraph.add_run(part)

def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def add_table(rows):
    # rows: list of list of cell strings; first row is header
    if not rows:
        return
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(cell_text.strip())
            run.font.size = Pt(10)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade_cell(cell, '0D9488')
            elif i % 2 == 0:
                shade_cell(cell, 'E9F7F4')
    doc.add_paragraph()

def flush_table(buf):
    rows = []
    for line in buf:
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        # skip separator row like |---|---|
        if all(set(c) <= set('-: ') for c in cells):
            continue
        rows.append(cells)
    if rows:
        add_table(rows)

lines = open(r'C:/Users/sudha/infertilityapp/Lyra_App_Report.md', encoding='utf-8').read().splitlines()

table_buf = []
in_code = False
code_buf = []

for line in lines:
    raw = line

    if raw.strip().startswith('```'):
        if in_code:
            # flush code block
            p = doc.add_paragraph()
            for cl in code_buf:
                run = p.add_run(cl + '\n')
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
            p.paragraph_format.left_indent = Inches(0.3)
            code_buf = []
            in_code = False
        else:
            flush_table(table_buf); table_buf = []
            in_code = True
        continue

    if in_code:
        code_buf.append(raw)
        continue

    # table accumulation
    if raw.strip().startswith('|'):
        table_buf.append(raw)
        continue
    else:
        flush_table(table_buf); table_buf = []

    stripped = raw.strip()
    if not stripped:
        continue

    if stripped == '---':
        doc.add_page_break()
        continue

    m = re.match(r'^(#{1,4})\s+(.*)', stripped)
    if m:
        level = len(m.group(1))
        text = m.group(2)
        doc.add_heading(text, level=min(level, 3))
        continue

    if re.match(r'^\d+\.\s', stripped):
        p = doc.add_paragraph(style='List Number')
        add_rich(p, re.sub(r'^\d+\.\s', '', stripped))
        continue

    if stripped.startswith('* ') or stripped.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        add_rich(p, stripped[2:])
        continue

    p = doc.add_paragraph()
    add_rich(p, stripped)

flush_table(table_buf)

# Title page header tweak: make first heading centered
if doc.paragraphs and doc.paragraphs[0].style.name.startswith('Heading'):
    doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

out = r'C:/Users/sudha/infertilityapp/Lyra_App_Report.docx'
doc.save(out)
print('saved:', out)
